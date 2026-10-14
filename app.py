import tkinter as tk
from tkinter import messagebox, filedialog, ttk
import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.drawing.image import Image as ExcelImage
import os
import time
import threading

from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By

def scegli_cartella():
    cartella = filedialog.askdirectory(title="Scegli dove salvare il file Excel")
    if cartella:
        var_cartella.set(cartella)

def avvia_analisi():
    url = entry_url.get().strip()
    cartella_salvataggio = var_cartella.get().strip()

    if not url:
        messagebox.showwarning("Attenzione", "Per favore, inserisci un link (URL) valido!")
        return
    if not cartella_salvataggio:
        messagebox.showwarning("Attenzione", "Per favore, scegli una cartella dove salvare il file!")
        return
    if not os.path.exists(cartella_salvataggio):
        messagebox.showwarning("Attenzione", "La cartella selezionata non esiste!")
        return

    bottone_analizza.config(state=tk.DISABLED)
    bottone_cartella.config(state=tk.DISABLED)
    barra_progresso['value'] = 0
    label_status.config(text="Avvio in corso... attendi")

    thread = threading.Thread(target=esegui_analisi_background, args=(url, cartella_salvataggio))
    thread.start()

def esegui_analisi_background(url, cartella_salvataggio):
    try:
        label_status.config(text="Apertura browser robot invisibile...")
        
        chrome_options = Options()
        chrome_options.add_argument("--headless")
        # Ingrandiamo un po' la finestra per vedere meglio gli elementi
        chrome_options.add_argument("--window-size=1920,1080")
        
        driver = webdriver.Chrome(options=chrome_options)
        driver.get(url)
        
        label_status.config(text="Caricamento pagina iniziale... (Attendo 3 secondi)")
        time.sleep(3)
        
        # --- SCORRIMENTO FLUIDO (UMANO) ---
        label_status.config(text="Scorrimento morbido per risvegliare tutte le card...")
        altezza_totale = driver.execute_script("return document.body.scrollHeight")
        posizione_attuale = 0

        # Scende di poco alla volta (400 pixel) per far credere al sito che un umano stia leggendo
        while posizione_attuale < altezza_totale:
            driver.execute_script("window.scrollBy(0, 400);")
            posizione_attuale += 400
            time.sleep(0.3) # Pausa piccolissima ad ogni "colpo di rotellina"
            altezza_totale = driver.execute_script("return document.body.scrollHeight")

        # Torniamo in cima
        driver.execute_script("window.scrollTo(0, 0);")
        time.sleep(1)
        # ----------------------------------

        wb = openpyxl.Workbook()
        giallo = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        
        if not os.path.exists("foto_temp"):
            os.makedirs("foto_temp")

        ws_cta = wb.active
        ws_cta.title = "CTA e Link"
        ws_cta.append(["Testo (Label)", "Link", "Screenshot (Foto)", "Note/Errori"])
        
        ws_cta.column_dimensions['A'].width = 30
        ws_cta.column_dimensions['B'].width = 50
        ws_cta.column_dimensions['C'].width = 40
        ws_cta.column_dimensions['D'].width = 30

        links = driver.find_elements(By.TAG_NAME, 'a')
        totale_links = len(links)
        
        immagini_web = driver.find_elements(By.TAG_NAME, 'img')
        totale_immagini = len(immagini_web)
        
        barra_progresso['maximum'] = totale_links + totale_immagini
        progresso_attuale = 0

        lista_cta = []

        for index, link in enumerate(links):
            progresso_attuale += 1
            barra_progresso['value'] = progresso_attuale
            label_status.config(text=f"Analisi Link {index + 1} di {totale_links}...")

            testo = "TESTO MANCANTE"
            href = "LINK MANCANTE"
            percorso_foto = None
            stato_screen = ""
            note_errore = ""
            is_giallo = False

            try:
                testo_estratto = link.text.strip()
                if not testo_estratto:
                    testo_estratto = link.get_attribute("innerText")
                    if testo_estratto:
                        testo_estratto = testo_estratto.strip()
                
                if testo_estratto:
                    testo = testo_estratto

                href_estratto = link.get_attribute('href')
                if href_estratto:
                    href = href_estratto

                if testo == "TESTO MANCANTE" or href == "LINK MANCANTE" or str(href).endswith("#"):
                    is_giallo = True

                # Centriamo il link sullo schermo PRIMA di controllare se è visibile
                driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", link)
                time.sleep(0.5) # Aspettiamo mezzo secondo per l'animazione di entrata

                if link.is_displayed() and link.size['width'] > 0 and link.size['height'] > 0:
                    percorso_foto = f"foto_temp/screen_{index}.png"
                    link.screenshot(percorso_foto)
                    stato_screen = "Foto Presente"
                else:
                    stato_screen = "Elemento Invisibile"
                    note_errore = "Nascosto o non visibile a schermo"

            except Exception as e:
                stato_screen = "Errore Fotocamera"
                note_errore = str(e).split('\n')[0]
                is_giallo = True

            lista_cta.append({
                "testo": testo,
                "href": href,
                "percorso_foto": percorso_foto,
                "stato_screen": stato_screen,
                "is_giallo": is_giallo,
                "note_errore": note_errore
            })

        # ORDINAMENTO
        def chiave_ordinamento(item):
            ha_foto = 0 if (item["percorso_foto"] and os.path.exists(item["percorso_foto"])) else 1
            e_giallo = 0 if item["is_giallo"] else 1
            return (e_giallo, ha_foto)

        lista_cta.sort(key=chiave_ordinamento)

        # SCRITTURA EXCEL
        for riga_index, item in enumerate(lista_cta, start=2):
            testo = item["testo"]
            href = item["href"]
            percorso_foto = item["percorso_foto"]
            stato_screen = item["stato_screen"]
            note_errore = item["note_errore"]
            is_giallo = item["is_giallo"]

            valore_colonna_c = "" if (percorso_foto and os.path.exists(percorso_foto)) else stato_screen
            
            ws_cta.append([testo, href, valore_colonna_c, note_errore])

            if testo == "TESTO MANCANTE":
                ws_cta.cell(row=riga_index, column=1).fill = giallo
            if href == "LINK MANCANTE" or str(href).endswith("#"):
                ws_cta.cell(row=riga_index, column=2).fill = giallo

            if percorso_foto and os.path.exists(percorso_foto):
                img = ExcelImage(percorso_foto)
                img.width = min(img.width, 250)
                img.height = min(img.height, 60)
                ws_cta.add_image(img, f"C{riga_index}")
                ws_cta.row_dimensions[riga_index].height = 50

        # FOGLIO 2: IMMAGINI
        ws_img = wb.create_sheet(title="Immagini e Alt Tag")
        ws_img.append(["Link Immagine", "Alt Tag"])
        ws_img.column_dimensions['A'].width = 80
        ws_img.column_dimensions['B'].width = 40

        riga_img = 2
        for index, img in enumerate(immagini_web):
            progresso_attuale += 1
            barra_progresso['value'] = progresso_attuale
            label_status.config(text=f"Analisi Immagine {index + 1} di {totale_immagini}...")

            try:
                src = img.get_attribute('src') or img.get_attribute('data-src') or "SORGENTE MANCANTE"
                alt = img.get_attribute('alt')
                if alt is None or alt.strip() == "":
                    alt = "ALT MANCANTE"
                    
                ws_img.append([src, alt])
                if alt == "ALT MANCANTE":
                    ws_img.cell(row=riga_img, column=2).fill = giallo
                riga_img += 1
            except Exception:
                pass

        driver.quit()

        nome_file = os.path.join(cartella_salvataggio, "Risultato_Analisi.xlsx")
        wb.save(nome_file)
        
        for f in os.listdir("foto_temp"):
            os.remove(os.path.join("foto_temp", f))
        os.rmdir("foto_temp")

        label_status.config(text="Analisi completata con successo!")
        messagebox.showinfo("Finito!", f"Tutto fatto!\nIl file è stato salvato in:\n{nome_file}")

    except Exception as e:
        label_status.config(text="Errore durante l'analisi.")
        messagebox.showerror("Errore", f"Si è verificato un problema:\n{e}")
    finally:
        bottone_analizza.config(state=tk.NORMAL)
        bottone_cartella.config(state=tk.NORMAL)

# --- INTERFACCIA GRAFICA ---
finestra = tk.Tk()
finestra.title("Analizzatore Visivo di Pagine Web")
finestra.geometry("600x420")
finestra.configure(bg="#f0f0f0")

tk.Label(finestra, text="1. Inserisci l'indirizzo della pagina web (URL):", font=("Arial", 11, "bold"), bg="#f0f0f0").pack(pady=(20, 5))
entry_url = tk.Entry(finestra, width=60, font=("Arial", 12))
entry_url.pack(pady=5)

tk.Label(finestra, text="2. Seleziona dove salvare il file:", font=("Arial", 11, "bold"), bg="#f0f0f0").pack(pady=(15, 5))
frame_cartella = tk.Frame(finestra, bg="#f0f0f0")
frame_cartella.pack(pady=5)

var_cartella = tk.StringVar()
tk.Entry(frame_cartella, textvariable=var_cartella, width=45, font=("Arial", 10), state="readonly").pack(side=tk.LEFT, padx=(0, 10))
bottone_cartella = tk.Button(frame_cartella, text="Sfoglia...", font=("Arial", 10), command=scegli_cartella)
bottone_cartella.pack(side=tk.LEFT)

bottone_analizza = tk.Button(finestra, text="Avvia Analisi", font=("Arial", 12, "bold"), bg="#0078D7", fg="white", padx=20, pady=5, command=avvia_analisi)
bottone_analizza.pack(pady=25)

barra_progresso = ttk.Progressbar(finestra, orient="horizontal", length=400, mode="determinate")
barra_progresso.pack(pady=5)

label_status = tk.Label(finestra, text="Pronto. Compila i campi qui sopra e clicca Avvia.", font=("Arial", 10, "italic"), bg="#f0f0f0")
label_status.pack(pady=5)

finestra.mainloop()